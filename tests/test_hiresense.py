import os
import io
import unittest
import tempfile
import sqlite3
from docx import Document

# Set test database path before importing app/database
TEST_DB = os.path.join(tempfile.gettempdir(), "test_hiresense.db")
os.environ["DATABASE_PATH"] = TEST_DB
os.environ["SECRET_KEY"] = "test-secret-key-123"

import database
from database import get_db_connection, init_db
from utils.parser import (
    extract_text, extract_contact_info, extract_text_from_txt,
    extract_text_from_docx, clean_name_from_filename
)
from utils.skills import extract_skills, parse_skill_list, compare_skills
from utils.scorer import calculate_text_similarity, evaluate_candidate
from utils.reporter import sanitize_csv_cell, generate_csv_report, generate_pdf_report
from utils.demo_data import seed_demo_data, clear_demo_data
import app as flask_app_module


class TestHireSenseCore(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        # Initialize test database
        if os.path.exists(TEST_DB):
            try:
                os.remove(TEST_DB)
            except OSError:
                pass
        database.DATABASE_PATH = TEST_DB
        init_db()

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(TEST_DB):
            try:
                os.remove(TEST_DB)
            except OSError:
                pass

    def setUp(self):
        self.app = flask_app_module.app
        self.app.config["TESTING"] = True
        self.app.config["UPLOAD_FOLDER"] = tempfile.gettempdir()
        self.client = self.app.test_client()

    # ==========================================
    # 1. PARSER & CONTACT EXTRACTION TESTS
    # ==========================================

    def test_txt_extraction(self):
        with tempfile.NamedTemporaryFile("w+", suffix=".txt", delete=False) as f:
            f.write("Aarav Sharma\nSoftware Engineer with Python and React skills.")
            txt_path = f.name
        
        try:
            text, err = extract_text(txt_path)
            self.assertIsNone(err)
            self.assertIn("Aarav Sharma", text)
            self.assertIn("Python", text)
        finally:
            if os.path.exists(txt_path):
                os.remove(txt_path)

    def test_docx_extraction(self):
        doc = Document()
        doc.add_heading("Priya Patel - Resume", 0)
        doc.add_paragraph("Full stack engineer with 3 years experience in Flask and SQL.")
        
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            docx_path = f.name
        
        try:
            doc.save(docx_path)
            text, err = extract_text(docx_path)
            self.assertIsNone(err)
            self.assertIn("Priya Patel", text)
            self.assertIn("Flask", text)
        finally:
            if os.path.exists(docx_path):
                os.remove(docx_path)

    def test_empty_and_corrupt_files(self):
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
            empty_path = f.name  # 0 bytes

        try:
            text, err = extract_text(empty_path)
            self.assertEqual(text, "")
            self.assertIn("empty", err.lower())
        finally:
            if os.path.exists(empty_path):
                os.remove(empty_path)

    def test_contact_info_extraction(self):
        sample_resume = (
            "Rohan Deshmukh\n"
            "Email: rohan.deshmukh@example.com | Phone: +91 98765 43210\n"
            "GitHub: https://github.com/rohandev\n"
            "LinkedIn: https://linkedin.com/in/rohandeshmukh\n"
            "Experience: 4 years in software engineering.\n"
            "Education: Bachelor of Technology (B.Tech) in Computer Science."
        )
        info = extract_contact_info(sample_resume, "Rohan_Resume.pdf")
        self.assertEqual(info["name"], "Rohan Deshmukh")
        self.assertEqual(info["email"], "rohan.deshmukh@example.com")
        self.assertIn("98765", info["phone"])
        self.assertIn("github.com/rohandev", info["github"])
        self.assertIn("linkedin.com/in/rohandeshmukh", info["linkedin"])
        self.assertEqual(info["years_experience"], 4.0)

    # ==========================================
    # 2. SKILL EXTRACTION & WORD-BOUNDARY TESTS
    # ==========================================

    def test_word_boundary_safety(self):
        # "academic", "experience", "computer" have the letter 'c', but NOT skill 'C'
        text_without_c_skill = "Academic experience in basic computer education and history."
        skills = extract_skills(text_without_c_skill)
        self.assertNotIn("C", skills, "Word boundary check failed: letter 'c' matched inside words!")

        # Explicit 'C' programming skill
        text_with_c = "Proficient in C programming and C++ development."
        skills_c = extract_skills(text_with_c)
        self.assertIn("C", skills_c)
        self.assertIn("C++", skills_c)

    def test_skill_synonyms(self):
        text = "Built web apps with React.js, Nodejs, and PostgreSQL database on AWS."
        skills = extract_skills(text)
        self.assertIn("React", skills)
        self.assertIn("Node.js", skills)
        self.assertIn("PostgreSQL", skills)
        self.assertIn("AWS", skills)

    def test_compare_skills(self):
        cand_skills = ["Python", "Flask", "SQL", "Git"]
        req_skills = ["Python", "Flask", "Docker"]
        pref_skills = ["Git", "AWS"]

        res = compare_skills(cand_skills, req_skills, pref_skills)
        self.assertIn("Python", res["matched_required"])
        self.assertIn("Flask", res["matched_required"])
        self.assertIn("Docker", res["missing_required"])
        self.assertIn("Git", res["matched_preferred"])
        self.assertIn("AWS", res["missing_preferred"])
        self.assertAlmostEqual(res["required_ratio"], 2/3, places=2)

    # ==========================================
    # 3. SCORER & EXPLAINABILITY TESTS
    # ==========================================

    def test_evaluate_candidate(self):
        resume_text = (
            "Senior Python Developer with 4 years experience.\n"
            "Expert in Python, Flask, React, and SQL. Deployed microservices using Git."
        )
        jd_text = (
            "Looking for a Python Developer with 3+ years experience.\n"
            "Must have Python, Flask, SQL, and Docker. Experience with Git is a plus."
        )
        req = ["Python", "Flask", "SQL", "Docker"]
        pref = ["Git"]

        eval_res = evaluate_candidate(
            resume_text=resume_text,
            job_description=jd_text,
            required_skills=req,
            preferred_skills=pref,
            min_experience_text="3+ years"
        )

        self.assertGreater(eval_res["overall_score"], 50.0)
        self.assertLessEqual(eval_res["overall_score"], 100.0)
        self.assertIn("Python", eval_res["matched_skills"])
        self.assertIn("Docker", eval_res["missing_skills"])
        self.assertTrue(len(eval_res["reasons"]) > 0)
        self.assertIn("Responsible AI", eval_res["disclaimer"])

    # ==========================================
    # 4. REPORTING & SECURITY TESTS
    # ==========================================

    def test_csv_injection_sanitization(self):
        # Malicious formulas in CSV cells must be prepended with a quote
        self.assertEqual(sanitize_csv_cell("=1+1"), "'=1+1")
        self.assertEqual(sanitize_csv_cell("+cmd|' /C calc'!A0"), "'+cmd|' /C calc'!A0")
        self.assertEqual(sanitize_csv_cell("@SUM(A1:A10)"), "'@SUM(A1:A10)")
        self.assertEqual(sanitize_csv_cell("-123"), "'-123")
        self.assertEqual(sanitize_csv_cell("Normal Text"), "Normal Text")

    def test_pdf_report_generation(self):
        dummy_cand = {
            "name": "Ananya Roy",
            "score": 88.5,
            "text_similarity": 82.0,
            "skill_score": 95.0,
            "status": "Shortlisted",
            "matched_skills": "Python, Flask, SQL",
            "missing_skills": "Docker",
            "notes": "Candidate passed initial screening with high marks."
        }
        dummy_job = {"title": "Backend Python Developer"}
        pdf_bytes = generate_pdf_report(dummy_cand, dummy_job)
        self.assertTrue(len(pdf_bytes) > 1000)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

    # ==========================================
    # 5. DEMO DATA SEED & CLEAR TESTS
    # ==========================================

    def test_demo_data_cycle(self):
        job_id = seed_demo_data()
        self.assertIsNotNone(job_id)

        conn = get_db_connection()
        demo_cnt = conn.execute("SELECT COUNT(*) FROM candidates WHERE is_demo = 1").fetchone()[0]
        self.assertGreater(demo_cnt, 0)
        conn.close()

        cleared_cnt = clear_demo_data()
        self.assertEqual(cleared_cnt, demo_cnt)

        conn = get_db_connection()
        after_clear_cnt = conn.execute("SELECT COUNT(*) FROM candidates WHERE is_demo = 1").fetchone()[0]
        self.assertEqual(after_clear_cnt, 0)
        conn.close()

    # ==========================================
    # 6. FLASK WEB ROUTES & ENDPOINTS
    # ==========================================

    def test_routes_status_codes(self):
        # Public routes
        for endpoint in ["/", "/dashboard", "/screen", "/jobs", "/pricing", "/about", "/settings", "/history", "/login", "/register"]:
            response = self.client.get(endpoint)
            self.assertIn(response.status_code, [200, 302], f"Failed on endpoint: {endpoint}")

    def test_recruiter_screening_flow(self):
        # Create a sample text resume
        data = {
            "job_selection": "new",
            "job_title": "Test AI Engineer",
            "department": "Engineering",
            "job_description": "We need an AI Engineer with Python, Flask, and Scikit-Learn skills.",
            "required_skills": "Python, Flask, Scikit-Learn",
            "preferred_skills": "Docker",
            "min_experience": "1-3 years (Junior)",
            "qualification": "Bachelor's Degree",
            "resumes": (io.BytesIO(b"Candidate Name: Test Applicant\nSkills: Python, Flask, Scikit-Learn, Git.\nEmail: test@example.com"), "Test_Applicant.txt")
        }

        res = self.client.post("/screen", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Test AI Engineer", res.data)
        self.assertIn(b"Test Applicant", res.data)

        # Check candidate was stored in DB
        conn = get_db_connection()
        cand = conn.execute("SELECT * FROM candidates WHERE name = 'Test Applicant'").fetchone()
        self.assertIsNotNone(cand)
        self.assertGreater(cand["score"], 50.0)
        
        # Test updating candidate status
        status_res = self.client.post(f"/candidate/{cand['id']}/status", data={"status": "Shortlisted"})
        self.assertIn(status_res.status_code, [200, 302])

        updated_cand = conn.execute("SELECT status FROM candidates WHERE id = ?", (cand["id"],)).fetchone()
        self.assertEqual(updated_cand["status"], "Shortlisted")

        # Test CSV export
        csv_res = self.client.get(f"/jobs/{cand['job_id']}/export/csv")
        self.assertEqual(csv_res.status_code, 200)
        self.assertIn(b"Test Applicant", csv_res.data)

        # Test PDF export
        pdf_res = self.client.get(f"/candidate/{cand['id']}/pdf")
        self.assertEqual(pdf_res.status_code, 200)
        self.assertTrue(pdf_res.data.startswith(b"%PDF"))

        conn.close()


if __name__ == "__main__":
    unittest.main()
