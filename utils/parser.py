import os
import re
import pdfplumber

try:
    import docx
except ImportError:
    docx = None

def extract_text_from_pdf(filepath):
    """
    Extract text from a PDF file using pdfplumber.
    Handles empty, password-protected, or corrupted files gracefully.
    """
    text_chunks = []
    try:
        with pdfplumber.open(filepath) as pdf:
            if not pdf.pages:
                return "", "PDF file contains no pages."
            for i, page in enumerate(pdf.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text_chunks.append(page_text.strip())
                except Exception as page_err:
                    continue
        full_text = "\n".join(text_chunks).strip()
        if not full_text:
            return "", "PDF document appears to be scanned or contains no extractable text."
        return full_text, None
    except Exception as e:
        err_msg = str(e)
        if "password" in err_msg.lower():
            return "", "PDF is password protected or encrypted."
        return "", f"Could not read PDF: {err_msg}"

def extract_text_from_docx(filepath):
    """
    Extract text from a DOCX file using python-docx.
    """
    if docx is None:
        return "", "DOCX parser library is not available."
    try:
        doc = docx.Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            return "", "DOCX document contains no readable text."
        return full_text, None
    except Exception as e:
        return "", f"Could not read DOCX: {str(e)}"

def extract_text_from_txt(filepath):
    """
    Extract text from a plain text file with encoding fallback.
    """
    encodings = ["utf-8", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                content = f.read().strip()
                if not content:
                    return "", "Text file is empty."
                return content, None
        except UnicodeDecodeError:
            continue
        except Exception as e:
            return "", f"Could not read file: {str(e)}"
    return "", "Unsupported file text encoding."

def extract_text(filepath):
    """
    Main extraction interface. Dispatches to the appropriate parser by file extension.
    Returns: (extracted_text: str, error_message: str or None)
    """
    if not os.path.exists(filepath):
        return "", "File does not exist."
    
    if os.path.getsize(filepath) == 0:
        return "", "Uploaded file is 0 bytes (empty file)."
    
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        if ext == ".doc":
            return "", "Legacy .doc format is deprecated; please convert to .docx or .pdf."
        return extract_text_from_docx(filepath)
    elif ext in [".txt", ".rtf"]:
        return extract_text_from_txt(filepath)
    else:
        return "", f"Unsupported file extension: '{ext}'. Please upload PDF, DOCX, or TXT."

def extract_contact_info(text, fallback_filename=""):
    """
    Extract contact metadata (name heuristic, email, phone, GitHub, LinkedIn, education, experience)
    using rule-based patterns and regex.
    """
    info = {
        "name": "",
        "email": "",
        "phone": "",
        "github": "",
        "linkedin": "",
        "education": [],
        "years_experience": None
    }
    
    if not text:
        info["name"] = clean_name_from_filename(fallback_filename)
        return info

    # 1. Email extraction
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b', text)
    if email_match:
        info["email"] = email_match.group(0).lower()

    # 2. Phone extraction (handles +91, dashes, spaces, parentheses)
    phone_match = re.search(r'(?:(?:\+|00)\d{1,3}[-.\s]?)?(?:\(?\d{2,5}\)?[-.\s]?)?\d{3,5}[-.\s]?\d{4,5}\b', text)
    if phone_match:
        phone_candidate = phone_match.group(0).strip()
        # Filter out common digit patterns like dates (2020-2024)
        if len(re.sub(r'\D', '', phone_candidate)) >= 10:
            info["phone"] = phone_candidate

    # 3. GitHub & LinkedIn
    gh_match = re.search(r'(?:https?:\/\/)?(?:www\.)?github\.com\/([A-Za-z0-9_-]+)', text, re.IGNORECASE)
    if gh_match:
        info["github"] = gh_match.group(0)
    elif "github" in text.lower():
        info["github"] = "Mentioned"

    li_match = re.search(r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([A-Za-z0-9_-]+)', text, re.IGNORECASE)
    if li_match:
        info["linkedin"] = li_match.group(0)
    elif "linkedin" in text.lower():
        info["linkedin"] = "Mentioned"

    # 4. Education degree detection
    edu_patterns = [
        r'\b(?:B\.?Tech|B\.?E\.?|Bachelor(?:\'s)?\s+of\s+Technology|Bachelor(?:\'s)?\s+of\s+Engineering)\b',
        r'\b(?:M\.?Tech|M\.?E\.?|Master(?:\'s)?\s+of\s+Technology|Master(?:\'s)?\s+of\s+Engineering)\b',
        r'\b(?:B\.?Sc|B\.?C\.?A\.?|Bachelor(?:\'s)?\s+of\s+Computer\s+Applications)\b',
        r'\b(?:M\.?Sc|M\.?C\.?A\.?|Master(?:\'s)?\s+of\s+Computer\s+Applications)\b',
        r'\b(?:Ph\.?D|Doctorate)\b',
        r'\b(?:Bachelor(?:\'s)?(?:\s+Degree)?|Master(?:\'s)?(?:\s+Degree)?|Diploma)\b'
    ]
    for pat in edu_patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            degree_str = m.group(0).strip()
            if degree_str not in info["education"]:
                info["education"].append(degree_str)

    # 5. Experience detection (e.g. "3 years", "5+ years of experience")
    exp_match = re.search(r'(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)(?:\s+(?:of\s+)?experience)?', text, re.IGNORECASE)
    if exp_match:
        try:
            info["years_experience"] = float(exp_match.group(1))
        except ValueError:
            pass

    # 6. Candidate Name heuristic
    # Check top lines of the resume text
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    candidate_name = ""
    for line in lines[:5]:
        # Strip common prefixes like 'Candidate Name:', 'Name:'
        line = re.sub(r'^(?:candidate\s+name|name|full\s+name|applicant\s+name)\s*:\s*', '', line, flags=re.IGNORECASE).strip()
        cleaned_line = re.sub(r'[^A-Za-z\s]', '', line).strip()
        words = cleaned_line.split()
        if 2 <= len(words) <= 4:
            lower_line = line.lower()
            if not any(header in lower_line for header in ["curriculum", "vitae", "resume", "summary", "profile", "contact", "experience", "education", "skills"]):
                candidate_name = cleaned_line.title()
                break
    
    if not candidate_name:
        candidate_name = clean_name_from_filename(fallback_filename)
        
    info["name"] = candidate_name
    return info

def clean_name_from_filename(filename):
    """Extract a reasonable human name from a resume filename."""
    if not filename:
        return "Unknown Candidate"
    base = os.path.splitext(os.path.basename(filename))[0]
    # Remove common tags
    base = re.sub(r'(?i)[_-]?(?:resume|cv|profile|latest|updated|\d+)', '', base)
    # Replace separators with spaces
    base = re.sub(r'[\._-]+', ' ', base).strip()
    return base.title() if base else "Candidate"